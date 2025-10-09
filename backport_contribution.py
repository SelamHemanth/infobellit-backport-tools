#!/usr/bin/env python3
"""
Git Backport Contribution Checker
Checks signed-off commits by email IDs across specified branches.
Each commit is assigned only to the first matching email (in user-provided order).
"""

import subprocess
import re
import pandas as pd
from datetime import datetime


def validate_email(email):
    """Basic email validation"""
    pattern = r'^[a-zA-Z0-9._%+-]+@[a-zA-Z0-9.-]+\.[a-zA-Z]{2,}$'
    return re.match(pattern, email) is not None


def check_branch_exists(branch):
    """Check if a branch exists in the repository"""
    try:
        subprocess.run(["git", "rev-parse", "--verify", branch],
                       capture_output=True, check=True)
        return True
    except subprocess.CalledProcessError:
        return False


def get_signedoff_commits(branch, since_date=None, until_date=None):
    """
    Get all commits in a branch with their signed-off emails.
    Returns a dict: {commit_hash: [list_of_signedoff_emails]}
    """
    try:
        cmd = ["git", "log", branch, "--format=%H%n%B%n---END---"]

        # Add date filters properly
        if since_date:
            cmd.append(f'--since={since_date}')
        if until_date and until_date.lower() != "now":
            cmd.append(f'--until={until_date}')
        elif until_date and until_date.lower() == "now":
            cmd.append(f'--until={datetime.now().isoformat()}')

        result = subprocess.run(
            cmd,
            capture_output=True,
            text=True,
            encoding='utf-8',
            errors='replace',
            check=True
        )

        output = result.stdout.strip()
        if not output:
            return {}

        commits = {}
        raw_commits = output.split("---END---")
        for entry in raw_commits:
            lines = entry.strip().splitlines()
            if not lines:
                continue
            commit_hash = lines[0].strip()
            message = "\n".join(lines[1:])
            signed_offs = re.findall(r"Signed-off-by:\s+.*<([^>]+)>", message, re.IGNORECASE)
            commits[commit_hash] = signed_offs
        return commits
    except subprocess.CalledProcessError as e:
        print(f"Error reading branch {branch}: {e}")
        return {}


def count_commits_per_email(emails, branch_commits):
    """
    Assign each commit to the first matching email in `emails` list.
    Returns a dict: {email: count}
    """
    counts = {email: 0 for email in emails}

    for commit_hash, signed_emails in branch_commits.items():
        signed_emails_lower = [e.lower() for e in signed_emails]
        for email in emails:
            if email.lower() in signed_emails_lower:
                counts[email] += 1
                break  # assign to first match only
    return counts


def main():
    print("=" * 60)
    print("Git Backport Contribution Checker")
    print("=" * 60)
    print()

    # Check if we're in a git repository
    try:
        subprocess.run(['git', 'rev-parse', '--git-dir'],
                       capture_output=True, check=True)
    except subprocess.CalledProcessError:
        print("Error: Not in a git repository!")
        return

    # Get email IDs
    print("Paste all email IDs (one per line or space-separated).")
    print("Type 'done' on a new line when finished:")
    print()

    emails = []
    while True:
        line = input().strip()
        if line.lower() == 'done':
            break
        if line:
            parts = re.split(r'[\s,]+', line)
            for email in parts:
                email = email.strip()
                if email and validate_email(email):
                    emails.append(email)
                    print(f"  ✓ Added: {email}")
                elif email:
                    print(f"  ✗ Invalid email: {email}")

    if not emails:
        print("No valid emails provided. Exiting.")
        return

    print(f"\n✓ Total emails added: {len(emails)}")
    print()

    # Get date range (optional)
    print("Enter date range for commit search (optional - press Enter to skip):")
    since_date = input("Start date (YYYY-MM-DD): ").strip()
    until_date = input("End date (YYYY-MM-DD or 'now'): ").strip()

    if since_date or until_date:
        print(f"  ✓ Date range: {since_date or 'beginning'} to {until_date or 'now'}")
    else:
        print("  ✓ No date filter applied (all commits)")
        since_date = None
        until_date = None
    print()

    # Get branches
    print("Paste all branch names (one per line or space-separated).")
    print("Type 'done' on a new line when finished:")
    print()

    branches = []
    while True:
        line = input().strip()
        if line.lower() == 'done':
            break
        if line:
            parts = re.split(r'[\s,]+', line)
            for branch in parts:
                branch = branch.strip()
                if branch:
                    if check_branch_exists(branch):
                        branches.append(branch)
                        print(f"  ✓ Branch found: {branch}")
                    else:
                        print(f"  ✗ Branch '{branch}' not found")

    if not branches:
        print("No valid branches provided. Exiting.")
        return

    print(f"\n✓ Total branches added: {len(branches)}")
    print()

    # Confirm submission
    print("=" * 60)
    print("Summary:")
    print(f"  Emails: {len(emails)}")
    print(f"  Branches: {len(branches)}")
    if since_date or until_date:
        print(f"  Date range: {since_date or 'beginning'} to {until_date or 'now'}")
    else:
        print(f"  Date range: All commits (no filter)")
    confirm = input("Proceed with analysis? (y/n): ").strip().lower()

    if confirm != 'y':
        print("Analysis cancelled.")
        return

    print("\n" + "=" * 60)
    print("Processing commits...")
    print("=" * 60)

    # Collect data efficiently
    data = []

    for branch in branches:
        print(f"\nAnalyzing branch: {branch}...")
        branch_commits = get_signedoff_commits(branch, since_date, until_date)
        branch_counts = count_commits_per_email(emails, branch_commits)

        for email in emails:
            row = next((r for r in data if r['Mail ID'] == email), None)
            if not row:
                row = {'Mail ID': email}
                data.append(row)
            row[branch] = branch_counts[email]

    # Compute totals
    for row in data:
        total = sum(row.get(branch, 0) for branch in branches)
        row['Total commits'] = total

    # Create DataFrame
    df = pd.DataFrame(data)
    column_order = ['Mail ID'] + branches + ['Total commits']
    df = df[column_order]

    # Transpose for Excel output
    df_transposed = df.set_index('Mail ID').T

    # Save file
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    filename = f"backport_contribution_{timestamp}.xlsx"

    from openpyxl import load_workbook
    from openpyxl.styles import PatternFill, Font, Alignment, Border, Side

    df_transposed.to_excel(filename, engine='openpyxl')
    wb = load_workbook(filename)
    ws = wb.active

    # Style definitions
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    total_fill = PatternFill(start_color="FFC000", end_color="FFC000", fill_type="solid")
    branch_fill = PatternFill(start_color="E7E6E6", end_color="E7E6E6", fill_type="solid")

    header_font = Font(bold=True, color="FFFFFF", size=11)
    total_font = Font(bold=True, size=11)
    branch_font = Font(size=10)
    center_align = Alignment(horizontal="center", vertical="center")
    left_align = Alignment(horizontal="left", vertical="center")
    thin_border = Border(
        left=Side(style='thin'), right=Side(style='thin'),
        top=Side(style='thin'), bottom=Side(style='thin')
    )

    # Header
    for col in range(1, ws.max_column + 1):
        cell = ws.cell(row=1, column=col)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = center_align
        cell.border = thin_border

    # Branch names (first column)
    for row in range(2, ws.max_row + 1):
        cell = ws.cell(row=row, column=1)
        cell.fill = branch_fill
        cell.font = branch_font
        cell.alignment = left_align
        cell.border = thin_border

    # Total row
    total_row = ws.max_row
    for col in range(1, ws.max_column + 1):
        cell = ws.cell(row=total_row, column=col)
        cell.fill = total_fill
        cell.font = total_font
        cell.alignment = center_align if col > 1 else left_align
        cell.border = thin_border

    # Data cells
    for row in range(2, ws.max_row):
        for col in range(2, ws.max_column + 1):
            cell = ws.cell(row=row, column=col)
            cell.alignment = center_align
            cell.border = thin_border

    # Auto column width
    for column in ws.columns:
        max_length = 0
        column_letter = column[0].column_letter
        for cell in column:
            if cell.value:
                max_length = max(max_length, len(str(cell.value)))
        ws.column_dimensions[column_letter].width = min(max_length + 2, 50)

    ws.column_dimensions['A'].width = max(ws.column_dimensions['A'].width, 35)
    wb.save(filename)

    print("\n" + "=" * 60)
    print("Analysis Complete!")
    print("=" * 60)
    print(f"\n✓ Report saved to: {filename}")
    print("\nSummary:")
    print(df.to_string(index=False))
    print()


if __name__ == "__main__":
    try:
        main()
    except KeyboardInterrupt:
        print("\n\nProcess interrupted by user.")
    except Exception as e:
        print(f"\nError: {e}")

